"""
Excel (.xlsx) exports for the Arabic reports.

A .csv cannot carry text direction, column widths or fonts, so Excel lays an
Arabic sheet out left-to-right with clipped columns and overflowing dates. The
spreadsheet export is therefore a real workbook: a right-to-left sheet, columns
sized to their content and a bold header row, built with openpyxl. The wording
comes from arabic.py, shared with the PDF; values stay as real numbers so the
sheet is still sortable and filterable.
"""
from datetime import date, datetime

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .arabic import CATEGORIES, LABELS, PRIORITIES, daily_rows

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

_HEADER_FILL = PatternFill("solid", fgColor="203E72")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14)
_RIGHT = Alignment(horizontal="right", vertical="center")


def _cell(value):
    """Dates as plain YYYY-MM-DD text; everything else as-is (real numbers)."""
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    return value


def _style_header_row(ws, row_index):
    for cell in ws[row_index]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT


def _finish(ws):
    """Right-to-left sheet, right-aligned cells, and columns sized to fit."""
    ws.sheet_view.rightToLeft = True
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = _RIGHT
    for column_cells in ws.columns:
        longest = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=0,
        )
        letter = get_column_letter(column_cells[0].column)
        ws.column_dimensions[letter].width = min(max(longest + 4, 14), 60)


def _response(wb, filename):
    response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def daily_xlsx(report, filename):
    """The daily officer report as a styled, right-to-left workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "التقرير اليومي"

    ws.append([LABELS["metric"], LABELS["value"]])
    for label, value in daily_rows(report):
        ws.append([label, _cell(value)])

    _style_header_row(ws, 1)
    _finish(ws)
    return _response(wb, filename)


def weekly_xlsx(report, filename):
    """The weekly summary as a styled, right-to-left workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "التقرير الأسبوعي"

    ws.append([LABELS["weekly_title"]])
    ws.append([LABELS["start_date"], _cell(report["start_date"])])
    ws.append([LABELS["end_date"], _cell(report["end_date"])])
    ws.append([])

    ws.append([LABELS["missions_by_priority"]])
    priority_header = ws.max_row + 1
    ws.append([LABELS["priority"], LABELS["count"]])
    for priority, count in report["missions_by_priority"].items():
        ws.append([PRIORITIES[priority], count])
    ws.append([])

    ws.append([LABELS["missions_by_category"]])
    category_header = ws.max_row + 1
    ws.append([LABELS["category"], LABELS["count"]])
    for category, count in report["missions_by_category"].items():
        ws.append([CATEGORIES[category], count])
    ws.append([])

    ws.append([
        f'{LABELS["avg_ack"]} ({LABELS["seconds"]})',
        report["average_acknowledgement_seconds"],
    ])
    ws.append([
        f'{LABELS["avg_completion"]} ({LABELS["seconds"]})',
        report["average_completion_seconds"],
    ])
    ws.append([])

    ws.append([LABELS["top_officers"]])
    top_header = ws.max_row + 1
    ws.append([
        LABELS["officer_id"],
        LABELS["officer"],
        LABELS["completed_missions"],
    ])
    for officer in report["top_officers"]:
        ws.append([
            officer["officer_id"],
            officer["officer_name"],
            officer["completed_missions"],
        ])

    for row_index in (priority_header, category_header, top_header):
        _style_header_row(ws, row_index)
    ws["A1"].font = _TITLE_FONT
    _finish(ws)
    return _response(wb, filename)
